#!/usr/bin/env python3
"""
Importer Script: Migrasi Data Mentoring & Adab dari Google Spreadsheet ke MariaDB
Target Database: library_v4 (via docker exec ais_absensi_db)
"""
import os
import openpyxl
import subprocess
import json
import re
import datetime

DB_USER = os.getenv('DB_USER', 'absensi_user')
DB_PASS = os.getenv('DB_PASSWORD', '')
DB_NAME = os.getenv('DB_NAME', 'library_v4')

def run_mysql(sql):
    cmd = ['docker', 'exec', '-i', 'ais_absensi_db', 'mysql', '-u', DB_USER, '-pabsensi_secret_pass_2026', DB_NAME, '--default-character-set=utf8mb4']
    proc = subprocess.Popen(
        cmd,
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True
    )
    stdout, stderr = proc.communicate(input=sql)
    if proc.returncode != 0:
        raise Exception(f"MySQL Error: {stderr}\nSQL snippet: {sql[:300]}")
    return stdout

def esc(val):
    if val is None:
        return "NULL"
    if isinstance(val, (int, float)):
        return str(val)
    if isinstance(val, (datetime.datetime, datetime.date)):
        return f"'{val.strftime('%Y-%m-%d %H:%M:%S')}'"
    s = str(val).replace("\\", "\\\\").replace("'", "\\'").replace("\r", "")
    return f"'{s}'"

def main():
    print("=== MEMULAI MIGRASI DATA SPREADSHEET KE MARIADB ===")
    xlsx_path = '/tmp/mentoring_fresh.xlsx' if os.path.exists('/tmp/mentoring_fresh.xlsx') else '/tmp/mentoring_siswa.xlsx'
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    # 1. AMBIL DATA SISWA DARI DB
    res = run_mysql("SELECT id, name, nis FROM students;")
    db_students_by_nis = {}
    db_students_by_name = {}
    for line in res.strip().split('\n')[1:]:
        parts = line.split('\t')
        if len(parts) >= 3:
            sid, sname, snis = int(parts[0]), parts[1].strip(), parts[2].strip()
            db_students_by_nis[snis] = sid
            db_students_by_name[sname.lower()] = sid

    print(f"Loaded {len(db_students_by_nis)} students from MariaDB.")

    # 2. AMBIL / SEED USERS GURU DARI SHEET 'users' & 'siswa'
    res_u = run_mysql("SELECT id, username, name FROM users;")
    db_users_by_name = {}
    for line in res_u.strip().split('\n')[1:]:
        parts = line.split('\t')
        if len(parts) >= 3:
            uid, uname, name = int(parts[0]), parts[1].strip(), parts[2].strip()
            db_users_by_name[name.lower()] = uid
            db_users_by_name[uname.lower()] = uid

    # Kumpulkan seluruh nama mentor
    ws_siswa = wb['siswa']
    mentors_to_ensure = set()
    for row in ws_siswa.iter_rows(min_row=2, values_only=True):
        if row[3]:
            mentors_to_ensure.add(row[3].strip())

    ws_users = wb['users']
    for row in ws_users.iter_rows(min_row=2, values_only=True):
        if row[0]:
            mentors_to_ensure.add(row[0].strip())

    # Default bcrypt password untuk akun guru yang belum ada
    default_pw_hash = '$2y$10$92IXUNpkjO0rOQ5byMi.Ye4oKoEa3Ro9llC/.og/at2.uheWG/igi' # 'password'

    user_insert_sql = []
    for m in mentors_to_ensure:
        clean_m = m.strip()
        m_lower = clean_m.lower()
        # Cek apakah ada kecocokan
        found_id = None
        for uname, uid in db_users_by_name.items():
            if m_lower == uname or (len(m_lower) > 5 and m_lower in uname) or (len(uname) > 5 and uname in m_lower):
                found_id = uid
                break
        if not found_id:
            # Generate username dari nama
            uname_candidate = re.sub(r'[^a-zA-Z0-9]', '', clean_m.split(',')[0].lower())
            user_insert_sql.append(
                f"INSERT INTO users (username, password, name, role) VALUES ('{uname_candidate}', '{default_pw_hash}', {esc(clean_m)}, 'staff');"
            )

    if user_insert_sql:
        print(f"Creating {len(user_insert_sql)} missing mentor user accounts...")
        run_mysql("\n".join(user_insert_sql))
        # Reload db_users
        res_u = run_mysql("SELECT id, username, name FROM users;")
        db_users_by_name = {}
        for line in res_u.strip().split('\n')[1:]:
            parts = line.split('\t')
            if len(parts) >= 3:
                uid, uname, name = int(parts[0]), parts[1].strip(), parts[2].strip()
                db_users_by_name[name.lower()] = uid
                db_users_by_name[uname.lower()] = uid

    def get_mentor_id(mentor_str):
        if not mentor_str:
            return 1
        m_lower = mentor_str.strip().lower()
        # Direct
        if m_lower in db_users_by_name:
            return db_users_by_name[m_lower]
        # Partial
        m_words = re.sub(r'[^a-zA-Z0-9 ]', '', m_lower).split()
        if m_words:
            first_word = m_words[0]
            for name, uid in db_users_by_name.items():
                if first_word in name:
                    return uid
        return 1 # Fallback admin

    # 3. MIGRASI MASTER KURIKULUM ADAB ('materi')
    ws_materi = wb['materi']
    curriculum_sql = ["TRUNCATE TABLE mentoring_curriculums;"]
    curriculums_imported = 0
    curriculum_map = {} # (grade_level, week_code) -> id

    for row in ws_materi.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        title = row[0].strip()
        grade_level = int(row[1]) if row[1] else 7
        week_code = str(row[2]).strip() if row[2] else 'Pekan 1'
        desc = row[3] or ''
        link = row[4] or ''
        t1, t2, t3, t4 = row[5], row[6], row[7], row[8]
        r1, r2, r3 = row[9], row[10], row[11]
        r4 = row[12] if len(row) > 12 else None

        curriculum_sql.append(
            f"INSERT INTO mentoring_curriculums (title, grade_level, week_code, description, file_link, task_1, task_2, task_3, task_4, reflection_prompt_1, reflection_prompt_2, reflection_prompt_3, reflection_prompt_4) "
            f"VALUES ({esc(title)}, {grade_level}, {esc(week_code)}, {esc(desc)}, {esc(link)}, {esc(t1)}, {esc(t2)}, {esc(t3)}, {esc(t4)}, {esc(r1)}, {esc(r2)}, {esc(r3)}, {esc(r4)});"
        )
        curriculums_imported += 1

    run_mysql("\n".join(curriculum_sql))
    print(f"✓ Berhasil mengimpor {curriculums_imported} modul materi ke 'mentoring_curriculums'.")

    # Ambil ID kurikulum untuk foreign key
    res_c = run_mysql("SELECT id, grade_level, week_code FROM mentoring_curriculums;")
    for line in res_c.strip().split('\n')[1:]:
        parts = line.split('\t')
        if len(parts) >= 3:
            cid, gr, wk = int(parts[0]), int(parts[1]), parts[2].strip().lower()
            curriculum_map[(gr, wk)] = cid

    # 4. MIGRASI PENUGASAN MENTOR-SANTRI ('siswa')
    assignments_sql = ["TRUNCATE TABLE mentoring_assignments;"]
    assignments_imported = 0
    student_mentor_map = {} # student_id -> mentor_id

    for row in ws_siswa.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        sname = str(row[0]).strip()
        nis_raw = str(row[1]).split('.')[0].strip()
        mentor_name = str(row[3]).strip() if row[3] else ''

        student_id = db_students_by_nis.get(nis_raw) or db_students_by_name.get(sname.lower())
        if not student_id:
            # Cari substring nama
            for k_name, sid in db_students_by_name.items():
                if sname.lower() in k_name or k_name in sname.lower():
                    student_id = sid
                    break

        if student_id:
            mentor_id = get_mentor_id(mentor_name)
            student_mentor_map[student_id] = mentor_id
            assignments_sql.append(
                f"INSERT INTO mentoring_assignments (student_id, mentor_user_id) VALUES ({student_id}, {mentor_id});"
            )
            assignments_imported += 1

    run_mysql("\n".join(assignments_sql))
    print(f"✓ Berhasil mengimpor {assignments_imported} relasi santri-mentor ke 'mentoring_assignments'.")

    # 5. MIGRASI JURNAL SESI MENTORING ('Master')
    ws_master = wb['Master']
    sessions_sql = ["TRUNCATE TABLE mentoring_sessions;"]
    sessions_imported = 0

    def parse_score(val):
        if val is None:
            return "NULL"
        try:
            f = float(val)
            return str(round(f, 1))
        except:
            return "NULL"

    def parse_notes(val):
        if val is None or str(val).strip() == '-' or str(val).strip() == '':
            return None
        return str(val).strip()

    for row in ws_master.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        
        # Deteksi format baris
        r1_str = str(row[1]).strip()
        if '/' in r1_str:
            # Format Baru (September onwards): [Timestamp, TahunAjaran, Tanggal, Mentee, Mentor, Sesi, ...]
            t_stamp = row[2] if isinstance(row[2], (datetime.datetime, datetime.date)) else row[0]
            mentee_name = str(row[3]).strip() if row[3] else ''
            mentor_name = str(row[4]).strip() if row[4] else ''
            session_code = str(row[5]).strip() if row[5] else 'P1-Selasa'
            month_name = t_stamp.strftime('%B') if isinstance(t_stamp, (datetime.datetime, datetime.date)) else 'September'
            year_val = t_stamp.year if isinstance(t_stamp, (datetime.datetime, datetime.date)) else 2026
        else:
            # Format Lama (Juli-Agustus): [Timestamp, Mentee, Mentor, Bulan, Tahun, Sesi, ...]
            t_stamp = row[0]
            mentee_name = r1_str
            mentor_name = str(row[2]).strip() if row[2] else ''
            month_name = str(row[3]).strip() if row[3] else 'Juli'
            try:
                year_val = int(float(row[4])) if row[4] else 2026
            except:
                year_val = 2026
            session_code = str(row[5]).strip() if row[5] else 'P1-Selasa'

        if not mentee_name:
            continue

        # Cari student_id
        student_id = db_students_by_name.get(mentee_name.lower())
        if not student_id:
            # Alias mapping untuk nama yang sedikit berbeda di sheet
            aliases = {
                "mikhail jovian hasbullah": "mikhail jovian",
                "hawa kami latuzzahra": "hawa kamilatuzzahra",
                "ayla jamayra ihsan": "ayla zamayra ihsan",
                "ahmad soleh, s.pd": "affan eferat hasan" # Kasus terbalik nama mentor & mentee
            }
            mapped_name = aliases.get(mentee_name.lower())
            if mapped_name and mapped_name in db_students_by_name:
                student_id = db_students_by_name[mapped_name]
            else:
                for k_name, sid in db_students_by_name.items():
                    if mentee_name.lower() in k_name or k_name in mentee_name.lower():
                        student_id = sid
                        break

        if not student_id:
            continue

        mentor_id = get_mentor_id(mentor_name)
        session_date = t_stamp.strftime('%Y-%m-%d') if isinstance(t_stamp, (datetime.datetime, datetime.date)) else '2026-08-01'

        score_mip = parse_score(row[6] if len(row) > 6 else None)
        notes_mip = parse_notes(row[7] if len(row) > 7 else None)

        score_acad = parse_score(row[8] if len(row) > 8 else None)
        notes_acad = parse_notes(row[9] if len(row) > 9 else None)

        score_mut = parse_score(row[10] if len(row) > 10 else None)
        notes_mut = parse_notes(row[11] if len(row) > 11 else None)

        score_adab = parse_score(row[12] if len(row) > 12 else None)
        notes_adab = parse_notes(row[13] if len(row) > 13 else None)

        score_fb = "NULL"
        notes_fb = parse_notes(row[14] if len(row) > 14 else None)

        sessions_sql.append(
            f"INSERT INTO mentoring_sessions (student_id, mentor_user_id, session_date, session_code, month_name, year_val, "
            f"score_mip, notes_mip, score_academic, notes_academic, score_mutabaah, notes_mutabaah, score_adab, notes_adab, score_feedback, notes_feedback, created_at) "
            f"VALUES ({student_id}, {mentor_id}, '{session_date}', {esc(session_code)}, {esc(month_name)}, {year_val}, "
            f"{score_mip}, {esc(notes_mip)}, {score_acad}, {esc(notes_acad)}, {score_mut}, {esc(notes_mut)}, {score_adab}, {esc(notes_adab)}, {score_fb}, {esc(notes_fb)}, {esc(t_stamp)});"
        )
        sessions_imported += 1
        if len(sessions_sql) >= 200:
            run_mysql("\n".join(sessions_sql))
            sessions_sql = []

    if sessions_sql:
        run_mysql("\n".join(sessions_sql))
    print(f"✓ Berhasil mengimpor {sessions_imported} sesi jurnal evaluasi ke 'mentoring_sessions'.")

    # 6. MIGRASI LOG CHECKLIST & REFLEKSI ADAB ('AdabSiswa')
    ws_adab = wb['AdabSiswa']
    adab_sql = ["TRUNCATE TABLE student_adab_logs;"]
    adab_imported = 0

    for row in ws_adab.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[1]:
            continue
        t_stamp = row[0]
        nis_raw = str(row[1]).split('.')[0].strip()
        student_name = str(row[2]).strip()
        kelas = str(row[3]).strip() if row[3] else '7A'
        pekan = str(row[4]).strip() if row[4] else 'Pekan 1'
        hari = str(row[5]).strip() if row[5] else 'Selasa'
        checklist_raw = str(row[6]).strip() if row[6] else '{}'
        r_action = row[7] or ''
        r_obstacle = row[8] or ''
        r_solution = row[9] or ''
        r_gratitude = row[10] or ''

        student_id = db_students_by_nis.get(nis_raw) or db_students_by_name.get(student_name.lower())
        if not student_id:
            for k_name, sid in db_students_by_name.items():
                if student_name.lower() in k_name or k_name in student_name.lower():
                    student_id = sid
                    break

        if not student_id:
            continue

        # Cari curriculum_id
        grade_num = 7
        if '8' in kelas: grade_num = 8
        elif '9' in kelas: grade_num = 9
        curr_id = curriculum_map.get((grade_num, pekan.lower()), "NULL")

        # Pastikan checklist_raw valid JSON
        try:
            json.loads(checklist_raw)
            valid_json = checklist_raw
        except:
            valid_json = "{}"

        adab_sql.append(
            f"INSERT INTO student_adab_logs (student_id, curriculum_id, grade_name, week_code, day_name, checklist_json, reflection_action, reflection_obstacle, reflection_solution, reflection_gratitude, submitted_at) "
            f"VALUES ({student_id}, {curr_id}, {esc(kelas)}, {esc(pekan)}, {esc(hari)}, {esc(valid_json)}, {esc(r_action)}, {esc(r_obstacle)}, {esc(r_solution)}, {esc(r_gratitude)}, {esc(t_stamp)});"
        )
        adab_imported += 1

        if len(adab_sql) >= 500:
            run_mysql("\n".join(adab_sql))
            adab_sql = []

    if adab_sql:
        run_mysql("\n".join(adab_sql))
    print(f"✓ Berhasil mengimpor {adab_imported} catatan amalan adab ke 'student_adab_logs'.")

    print("\n🎉 MIGRASI DATA LENGKAP SELESAI TANPA KENDALA!")

if __name__ == '__main__':
    main()
